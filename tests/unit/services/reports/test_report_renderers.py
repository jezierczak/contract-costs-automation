from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from contract_costs.services.reports.renders.cli import render_stdout
from contract_costs.services.reports.renders.excel import ExcelReportRenderer


def test_render_stdout_prints_no_data_for_empty_dataframe(capsys) -> None:
    render_stdout(pd.DataFrame())
    captured = capsys.readouterr()
    assert "No data." in captured.out


def test_excel_report_renderer_formats_numeric_and_marks_suma(tmp_path) -> None:
    df = pd.DataFrame(
        [
            {"contract_code": "C1", "net_amount": 100.0, "Wynik": 20.0},
            {"contract_code": "SUMA", "net_amount": 100.0, "Wynik": -5.0},
        ]
    )
    output = tmp_path / "report.xlsx"

    ExcelReportRenderer.render(df, output_path=output, sheet_name="report")

    wb = load_workbook(output)
    ws = wb["report"]
    assert ws["B2"].number_format == "#,##0.00"
    assert ws["A3"].value == "SUMA"
    assert ws["A3"].font.bold is True


def test_excel_report_renderer_excel_col_letter() -> None:
    assert ExcelReportRenderer._excel_col_letter(1) == "A"
    assert ExcelReportRenderer._excel_col_letter(27) == "AA"
    assert ExcelReportRenderer._excel_col_letter(52) == "AZ"

