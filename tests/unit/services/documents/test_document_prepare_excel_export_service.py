from datetime import date
from uuid import uuid4

from openpyxl import load_workbook

from contract_costs.services.documents.prepare.document_prepare_excel_export_service import (
    DocumentPrepareExcelExportService,
)
from contract_costs.services.documents.prepare.dto.candidate_record_dto import (
    CandidateRecordDto,
)
from contract_costs.services.documents.prepare.dto.prepare_document_bundle import (
    PrepareDocumentsBundle,
)
from contract_costs.services.documents.prepare.dto.prepare_document_dto import (
    PreparedDocumentDto,
)


def test_document_prepare_excel_export_service_exports_placeholder_dict_when_no_candidates(
    tmp_path,
) -> None:
    bundle = PrepareDocumentsBundle(documents=[])
    out = tmp_path / "documents.xlsx"

    DocumentPrepareExcelExportService.export(
        organization_id="org-1",
        bundle=bundle,
        output_path=out,
    )

    wb = load_workbook(out)
    ws = wb["_dict_existing_records"]
    assert ws["A2"].value == "NONE"


def test_document_prepare_excel_export_service_exports_candidate_dictionary(tmp_path) -> None:
    candidate = CandidateRecordDto(
        record_id=uuid4(),
        reference="FV/123",
        status="in_progress",
        invoice_date=date(2026, 2, 16),
    )
    document = PreparedDocumentDto(
        document_id=uuid4(),
        document_source="pdf",
        document_type="invoice",
        document_number="FV/123",
        seller_nip="1234567890",
        file_path="incoming/documents/fv_123.pdf",
        candidates=[candidate],
        confidence_score=None,
        confidence_breakdown=None,
    )
    bundle = PrepareDocumentsBundle(documents=[document])
    out = tmp_path / "documents_with_candidate.xlsx"

    DocumentPrepareExcelExportService.export(
        organization_id="org-1",
        bundle=bundle,
        output_path=out,
    )

    wb = load_workbook(out)
    ws = wb["_dict_existing_records"]
    assert ws["A2"].value == "1234567890"
    assert "FV/123" in str(ws["B2"].value)
    assert ws["C2"].value == str(candidate.record_id)

