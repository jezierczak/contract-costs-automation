from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook

from contract_costs.infrastructure.excel.excel_loader import ExcelLoader
from contract_costs.model.document import DocumentType
from contract_costs.services.documents.apply.dto.apply_document_command import DocumentApplyAction, ApplyDocumentCommand
from contract_costs.services.documents.prepare.build_document_assignment_columns import \
    build_document_assignment_columns


class DocumentActionExcelLoader:

    def load(
        self,
        *,
        excel_path: Path,
        organization_id: UUID,
        actor_user_id: UUID,
    ) -> list[ApplyDocumentCommand]:

        columns = build_document_assignment_columns()

        rows = ExcelLoader.load(
            input_path=excel_path,
            columns=columns,
            sheet_name="Documents",
            start_row=5,
        )

        record_label_map = self._load_record_dictionary(excel_path)

        commands: list[ApplyDocumentCommand] = []

        for row in rows:

            # -----------------------------
            # 1️⃣ BASIC EXTRACTION
            # -----------------------------

            if not row.get("Action"):
                continue

            try:
                action = DocumentApplyAction(row["Action"])
            except ValueError:
                raise ValueError(
                    f"Invalid action '{row['Action']}' "
                    f"for document {row.get('ID')}"
                )

            try:
                document_id = UUID(row["ID"])
            except Exception:
                raise ValueError(
                    f"Invalid document ID '{row.get('ID')}'"
                )

            target_record_id: UUID | None = None

            # -----------------------------
            # 2️⃣ ACTION LOGIC
            # -----------------------------

            if action == DocumentApplyAction.ADD_TO_EXISTING:

                label = row.get("Attach To Record")

                if not label:
                    raise ValueError(
                        f"Missing target record for document {document_id}"
                    )

                if label not in record_label_map:
                    raise ValueError(
                        f"Unknown record label '{label}' "
                        f"for document {document_id}"
                    )

                target_record_id = record_label_map[label]

            else:
                # dla innych akcji ignorujemy kolumnę Attach
                target_record_id = None

            # -----------------------------
            # 3️⃣ OVERRIDES
            # -----------------------------

            override_type = row.get("Type")
            if override_type:
                try:
                    DocumentType(override_type)
                except ValueError:
                    raise ValueError(f"Invalid document type '{override_type}'")

            override_number = row.get("Document Number")
            override_seller_nip = row.get("Seller NIP")

            if action in (
                DocumentApplyAction.SKIP,
                DocumentApplyAction.DELETE,
            ):
                # dla tych akcji nie ma sensu override
                override_type = None
                override_number = None
                override_seller_nip = None

            # -----------------------------
            # 4️⃣ BUILD COMMAND
            # -----------------------------

            cmd = ApplyDocumentCommand(
                organization_id=organization_id,
                actor_user_id=actor_user_id,
                document_id=document_id,
                action=action,
                target_record_id=target_record_id,
                override_document_type=override_type,
                override_document_number=override_number,
                override_seller_nip=override_seller_nip,
            )

            commands.append(cmd)

        return commands

    @staticmethod
    def _load_record_dictionary(

            excel_path: Path,
    ) -> dict[str, UUID]:

        wb = load_workbook(excel_path, data_only=True)

        if "_dict_existing_records" not in wb.sheetnames:
            return {}

        ws = wb["_dict_existing_records"]

        headers = {
            ws.cell(row=1, column=i).value: i
            for i in range(1, ws.max_column + 1)
        }

        label_col = headers["VALUE"]
        id_col = headers["RECORD_ID"]

        mapping: dict[str, UUID] = {}

        for row in range(2, ws.max_row + 1):

            label_raw = ws.cell(row=row, column=label_col).value
            record_id_raw = ws.cell(row=row, column=id_col).value

            if not label_raw or not record_id_raw:
                continue

            label = str(label_raw).strip()
            record_id_str = str(record_id_raw).strip()

            try:
                mapping[label] = UUID(record_id_str)
            except ValueError:
                continue

        return mapping
